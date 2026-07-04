#!/usr/bin/env python3
"""Rebuild aeora_prospects.xlsx from aeora_prospects.csv.

The CSV is the master file - edit/append rows there (or in any spreadsheet
app that saves back to CSV), then run:

    python3 build_xlsx.py

to regenerate a formatted Excel workbook with filters, frozen header,
priority colour-coding and a summary sheet.

Requires: pip install openpyxl
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(HERE, "aeora_prospects.csv")
XLSX_PATH = os.path.join(HERE, "aeora_prospects.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
HOT_FILL = PatternFill("solid", fgColor="C6EFCE")    # priority >= 8
WARM_FILL = PatternFill("solid", fgColor="FFF2CC")   # priority 7
WRAP = Alignment(wrap_text=True, vertical="top")

# Reasonable display widths per column index (1-based)
WIDTHS = [26, 22, 28, 30, 28, 16, 26, 30, 26, 46, 12, 12, 12, 12, 34, 40, 46, 10, 12, 34]


def main():
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        reader = list(csv.reader(f))
    header, data = reader[0], reader[1:]
    pri_idx = header.index("Priority score /10")

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"

    ws.append(header)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in data:
        typed = list(row)
        try:
            typed[pri_idx] = int(typed[pri_idx])
        except (ValueError, IndexError):
            pass
        ws.append(typed)

    for r in range(2, ws.max_row + 1):
        pri = ws.cell(row=r, column=pri_idx + 1).value
        fill = None
        if isinstance(pri, int):
            fill = HOT_FILL if pri >= 8 else WARM_FILL if pri == 7 else None
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = WRAP
            if fill:
                cell.fill = fill

    for i, w in enumerate(WIDTHS[: len(header)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    s = wb.create_sheet("Summary")
    s.append(["Aeora prospect database - summary"])
    s["A1"].font = Font(bold=True, size=12)
    s.append([])
    s.append(["Total prospects", len(data)])
    cats, locs = {}, {}
    for row in data:
        cats[row[1]] = cats.get(row[1], 0) + 1
    s.append([])
    s.append(["By category"])
    s["A5"].font = Font(bold=True)
    for cat, n in sorted(cats.items(), key=lambda kv: -kv[1]):
        s.append([cat, n])
    s.column_dimensions["A"].width = 40

    wb.save(XLSX_PATH)
    print(f"wrote {len(data)} prospects -> {XLSX_PATH}")


if __name__ == "__main__":
    main()
